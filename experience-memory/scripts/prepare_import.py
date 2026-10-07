"""只读 ZIP/Excel 生成待复核快照；不将表内任意数字当作成交价格。"""
import argparse
import hashlib
import io
import json
import zipfile
from datetime import date,datetime
from pathlib import Path


def safe(v):
    return v.isoformat() if isinstance(v,(date,datetime)) else v


def prepare(source, output):
    import openpyxl
    dest=Path(output);dest.mkdir(parents=True,exist_ok=True)
    manifest=[]
    with zipfile.ZipFile(source) as z:
        for entry in z.infolist():
            if entry.is_dir():continue
            if entry.file_size>100*1024*1024:raise ValueError('单文件过大，需独立处理')
            name=entry.filename
            if not entry.flag_bits & 0x800:
                try:name=name.encode('cp437').decode('gbk')
                except (UnicodeError,ValueError):pass
            raw=z.read(entry);sha=hashlib.sha256(raw).hexdigest()
            info={'source_file':name,'sha256':sha,'status':'pending_review','size':len(raw)}
            if name.lower().endswith(('.xlsx','.xlsm')):
                wf=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
                wv=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
                sheets=[]
                try:
                    for s in wf:
                        values=wv[s.title];rows=[]
                        for fr,vr in zip(s.iter_rows(),values.iter_rows()):
                            cells=[]
                            for fc,vc in zip(fr,vr):
                                if fc.value is not None:cells.append({'cell':fc.coordinate,'raw':safe(fc.value),'cached':safe(vc.value)})
                            if cells:rows.append(cells)
                        sheets.append({'sheet':s.title,'state':s.sheet_state,'rows':rows})
                finally:wf.close();wv.close()
                file=sha+'.json';(dest/file).write_text(json.dumps({'source':info,'sheets':sheets},ensure_ascii=False,indent=2),encoding='utf-8')
                info.update(kind='price_source_candidate',snapshot=file,sheets=len(sheets),nonempty_rows=sum(len(s['rows']) for s in sheets))
            elif name.lower().endswith('.md'):
                file=sha+'.md';(dest/file).write_bytes(raw);info.update(kind='reference_candidate',snapshot=file)
            else:info.update(kind='unsupported',note='未自动读取，交由相应文件能力处理')
            manifest.append(info)
    result={'status':'staged_not_imported','files':manifest,'cloud_status':'not_uploaded','note':'保留原值与来源；需 AI 结构化和复核，不是已确认历史价格库'}
    (dest/'manifest.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    return result


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--zip',required=True);p.add_argument('--out',required=True)
    a=p.parse_args();r=prepare(a.zip,a.out);print(json.dumps({'status':r['status'],'files':len(r['files'])},ensure_ascii=False))
